"""DOCX extraction that emits explicit Markdown-like structural markers."""

from pathlib import Path

from docx import Document

from app.parsers.base import DocumentExtractionError, ExtractedDocument


class DocxParser:
    def parse(self, path: Path) -> ExtractedDocument:
        try:
            document = Document(str(path))
        except Exception as exc:
            raise DocumentExtractionError(f"Unable to read DOCX: {exc}") from exc

        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style else "").lower()
            if style == "title":
                blocks.append(f"# {text}")
            elif style.startswith("heading"):
                try:
                    level = min(6, max(1, int(style.split()[-1])))
                except ValueError:
                    level = 2
                blocks.append(f"{'#' * level} {text}")
            elif "list" in style:
                blocks.append(f"- {text}")
            else:
                blocks.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))

        if not blocks:
            raise DocumentExtractionError("The uploaded DOCX contains no extractable text.")
        return ExtractedDocument(text="\n\n".join(blocks), metadata={"format": "docx"})
