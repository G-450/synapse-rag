"""Automated structural ingestion and hierarchical retrieval tests.

Run from ``python/`` with:
    python scripts/test_upload_and_chunking.py

The suite creates PDF, DOCX, and TXT fixtures in a temporary directory and uses
an in-memory Qdrant collection. It never reads or modifies the project database.
"""

import asyncio
import math
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document as WordDocument
from langchain_core.embeddings import Embeddings
from pypdf import PdfWriter
from qdrant_client import QdrantClient

from app.chunking import ParentChildSplitter
from app.parsers.base import DocumentExtractionError, extract_document
from app.parsers.legal_hierarchy_parser import LegalHierarchyParser
from app.routes.chat import _ai_sdk_stream
import app.rag as rag
import app.routes.upload as upload_route


CONTRACT = """ARTICLE IV - CONFIDENTIALITY
Section 4.1 Confidential Information
The Recipient shall protect all confidential technical, financial, customer, and product information using reasonable care. The Recipient may use confidential information only to evaluate the proposed transaction and shall not disclose it to any third party.

Section 4.2 Required Disclosure
The Recipient may disclose confidential information when required by applicable law, but only after providing prompt written notice and reasonable assistance to the Discloser.

ARTICLE IX - TERMINATION
Section 9.1 Termination Rights
Either party may terminate this Agreement upon thirty days written notice if the other party materially breaches and fails to cure that breach during the notice period.
"""


class KeywordEmbeddings(Embeddings):
    """Small deterministic embedder for isolated retrieval tests."""

    terms = ("confidential", "disclose", "terminate", "breach")

    def _embed(self, text: str) -> list[float]:
        lowered = text.lower()
        vector = [float(lowered.count(term)) for term in self.terms] + [0.0] * 380
        norm = math.sqrt(sum(value * value for value in vector)) or 1.0
        return [value / norm for value in vector]

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._embed(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._embed(text)


def write_text_pdf(path: Path, text: str) -> None:
    """Write a tiny standards-compliant, selectable-text PDF without test-only dependencies."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)").replace("\n", " ")
    stream = f"BT /F1 10 Tf 40 750 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    path.write_bytes(pdf)


class StructuralIngestionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.original_get_embeddings = rag.get_embeddings
        self.original_upload_get_vectorstore = upload_route.get_vectorstore
        self.original_upload_get_qdrant_client = upload_route.get_qdrant_client

    def tearDown(self) -> None:
        rag._qdrant_client = None
        rag.get_embeddings = self.original_get_embeddings
        upload_route.get_vectorstore = self.original_upload_get_vectorstore
        upload_route.get_qdrant_client = self.original_upload_get_qdrant_client
        self.temp_dir.cleanup()

    def test_pdf_docx_and_txt_extraction(self) -> None:
        txt_path = self.root / "contract.txt"
        txt_path.write_text(CONTRACT, encoding="utf-8")

        docx_path = self.root / "contract.docx"
        docx = WordDocument()
        docx.add_heading("ARTICLE IV - CONFIDENTIALITY", level=1)
        docx.add_heading("Section 4.1 Confidential Information", level=2)
        docx.add_paragraph("The Recipient shall protect confidential information.")
        docx.add_paragraph("Notice obligations", style="List Bullet")
        docx.save(docx_path)

        pdf_path = self.root / "contract.pdf"
        write_text_pdf(pdf_path, CONTRACT)

        extracted_documents = [
            extract_document(txt_path),
            extract_document(docx_path),
            extract_document(pdf_path),
        ]
        self.assertIn("Section 4.1", extracted_documents[0].text)
        self.assertIn("## Section 4.1", extracted_documents[1].text)
        self.assertIn("CONFIDENTIALITY", extracted_documents[2].text)
        for index, extracted in enumerate(extracted_documents):
            parents, children = ParentChildSplitter(child_tokens=20, child_overlap=4).split(
                extracted.text, f"format-{index}", f"contract-{index}"
            )
            self.assertTrue(parents)
            self.assertTrue(children)

        blank_pdf = self.root / "scanned.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with blank_pdf.open("wb") as target:
            writer.write(target)
        with self.assertRaisesRegex(DocumentExtractionError, "Scanned images are not yet supported"):
            extract_document(blank_pdf)

    def test_hierarchy_parent_child_integrity(self) -> None:
        sections = LegalHierarchyParser().parse(CONTRACT)
        self.assertTrue(any(section.number == "4.1" for section in sections))
        self.assertTrue(any(section.article == "Article IV" for section in sections))
        section = next(section for section in sections if section.number == "4.1")
        self.assertEqual(section.title, "Confidential Information")
        self.assertFalse(any(section.kind == "article" for section in sections))

        parents, children = ParentChildSplitter(child_tokens=16, child_overlap=4).split(
            CONTRACT, "document-1", "contract.txt"
        )
        parent_ids = {parent.id for parent in parents}
        self.assertGreater(len(children), len(parents))
        for child in children:
            self.assertIn(child.parent_id, parent_ids)
            self.assertIn("[Document: contract.txt", child.content)
            self.assertTrue(child.parent_content)
            self.assertLess(child.start_char, child.end_char)

    def test_qdrant_metadata_retrieval_and_parent_deduplication(self) -> None:
        rag._qdrant_client = QdrantClient(":memory:")
        rag.get_embeddings = lambda: KeywordEmbeddings()
        _, children = ParentChildSplitter(child_tokens=18, child_overlap=8).split(
            CONTRACT, "document-qdrant", "contract.txt"
        )
        inserted = upload_route.index_children(children, batch_size=2)
        self.assertEqual(inserted, len(children))

        records, _ = rag._qdrant_client.scroll(
            collection_name="synapse_rag", limit=100, with_payload=True, with_vectors=False
        )
        metadata = records[0].payload["metadata"]
        for key in ("document_id", "filename", "source_corpus", "parent_id", "parent_header", "section_number", "chunk_type", "content", "parent_content"):
            self.assertIn(key, metadata)

        results = rag.retrieve_chunks(
            "When may confidential information be disclosed?",
            document_id="document-qdrant",
            limit=5,
            rerank=False,
        )
        self.assertTrue(results)
        self.assertTrue(all(item["document_id"] == "document-qdrant" for item in results))
        duplicated = [results[0], dict(results[0])]
        context = rag.format_context(duplicated, max_parent_clauses=5)
        self.assertEqual(context.count("--- [Parent Clause"), 1)
        self.assertEqual(context.count(results[0]["parent_content"]), 1)

    def test_chat_uses_ai_sdk_v7_ui_message_stream(self) -> None:
        class Chunk:
            content = "Streamed answer"

        async def chunks():
            yield Chunk()

        async def collect() -> str:
            return "".join([part async for part in _ai_sdk_stream(chunks(), [{"section": "4.1"}])])

        stream = asyncio.run(collect())
        self.assertIn('"type": "start"', stream)
        self.assertIn('"type": "data-citations"', stream)
        self.assertIn('"type": "text-delta"', stream)
        self.assertIn('"delta": "Streamed answer"', stream)
        self.assertTrue(stream.endswith("data: [DONE]\n\n"))

    def test_failed_batch_rolls_back_all_uploaded_chunks(self) -> None:
        rag._qdrant_client = QdrantClient(":memory:")
        rag.get_embeddings = lambda: KeywordEmbeddings()
        real_store = rag.get_vectorstore()
        _, children = ParentChildSplitter(child_tokens=8, child_overlap=2).split(
            CONTRACT, "rollback-document", "rollback.txt"
        )

        class FailAfterFirstBatch:
            calls = 0

            def add_documents(self, documents, ids):
                self.calls += 1
                if self.calls > 1:
                    raise RuntimeError("simulated embedding failure")
                return real_store.add_documents(documents=documents, ids=ids)

        upload_route.get_vectorstore = lambda: FailAfterFirstBatch()
        upload_route.get_qdrant_client = lambda: rag._qdrant_client

        with self.assertRaisesRegex(RuntimeError, "simulated embedding failure"):
            upload_route.index_children(children, batch_size=2)

        records, _ = rag._qdrant_client.scroll(
            collection_name="synapse_rag", limit=100, with_payload=True, with_vectors=False
        )
        self.assertEqual(records, [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
